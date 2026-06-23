"""Helpers OXML bas niveau pour python-docx (operations non couvertes par l'API).

Champ de pagination Word, letter-spacing, ombrage et bordures de cellule,
sections pleine page (bleed) et controle des en-tetes/pieds.
"""
from __future__ import annotations

from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import RGBColor


def _sub(parent, tag):
    el = OxmlElement(tag)
    parent.append(el)
    return el


def add_page_number_field(paragraph) -> None:
    """Ajoute un champ Word { PAGE } (numero de page auto)."""
    run = paragraph.add_run()
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r.append(begin)
    r.append(instr)
    r.append(end)
    return run


def set_char_spacing(run, twips: int) -> None:
    """Letter-spacing (w:spacing) en 1/20 de point."""
    rpr = run._r.get_or_add_rPr()
    spacing = rpr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        rpr.append(spacing)
    spacing.set(qn("w:val"), str(int(twips)))


def set_complex_font(run, name: str) -> None:
    """Force aussi la police complex-script (w:cs) pour coherence."""
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:cs"), name)


def set_cell_shading(cell, hex_color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)


def set_cell_vertical_alignment(cell, value: str = "center") -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    va = tcpr.find(qn("w:vAlign"))
    if va is None:
        va = OxmlElement("w:vAlign")
        tcpr.append(va)
    va.set(qn("w:val"), value)


def set_table_borders(table, *, color="FFFFFF", size=4, val="single",
                      which=("top", "bottom", "insideH")) -> None:
    """Pose des bordures coherentes (par defaut : fines lignes blanches inter-rangs)."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        if edge in which:
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "none")
            el.set(qn("w:sz"), "0")
            el.set(qn("w:space"), "0")


def set_cell_margins(cell, *, top=40, bottom=40, left=90, right=90) -> None:
    """Marges internes de cellule en twips (1/20 pt)."""
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for name, val in (("top", top), ("bottom", bottom), ("start", left),
                      ("end", right), ("left", left), ("right", right)):
        el = mar.find(qn(f"w:{name}"))
        if el is None:
            el = OxmlElement(f"w:{name}")
            mar.append(el)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    """Repete la ligne d'en-tete sur chaque page (tblHeader)."""
    tr_pr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    tr_pr.append(th)


def make_section_fullbleed(section) -> None:
    """Section pleine page : marges nulles, distances en-tete/pied nulles."""
    from docx.shared import Mm
    section.top_margin = Mm(0)
    section.bottom_margin = Mm(0)
    section.left_margin = Mm(0)
    section.right_margin = Mm(0)
    section.gutter = Mm(0)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)


def unlink_headers_footers(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.first_page_header.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False


def set_no_spacing(paragraph) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = 0
    pf.space_after = 0
    pf.line_spacing = 1.0


_ANCHOR_XML = (
    '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'behindDoc="1" distT="0" distB="0" distL="0" distR="0" simplePos="0" '
    'locked="0" layoutInCell="1" allowOverlap="1" relativeHeight="0">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="{cx}" cy="{cy}"/>'
    '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapNone/>'
    '<wp:docPr id="{pid}" name="bg{pid}"/>'
    '<wp:cNvGraphicFramePr/>'
    '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
    '<pic:pic><pic:nvPicPr><pic:cNvPr id="{pid}" name="bg{pid}"/><pic:cNvPicPr/></pic:nvPicPr>'
    '<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
    '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
    '</a:graphicData></a:graphic></wp:anchor>'
)

_PID = [1000]


def add_fullpage_background(paragraph, image_path: str, width_emu: int,
                            height_emu: int) -> None:
    """Insere une image ancree pleine page << derriere le texte >> (bleed).

    Ne consomme aucune hauteur de flux : pas de page blanche parasite.
    """
    from docx.oxml import parse_xml

    run = paragraph.add_run()
    inline_shape = run.add_picture(image_path, width=width_emu, height=height_emu)
    inline = inline_shape._inline
    blip = inline.find(qn("a:blip"))
    if blip is None:
        blip = inline.find(".//" + qn("a:blip"))
    rid = blip.get(qn("r:embed"))

    drawing = run._r.find(qn("w:drawing"))
    _PID[0] += 1
    anchor = parse_xml(_ANCHOR_XML.format(cx=width_emu, cy=height_emu,
                                          rid=rid, pid=_PID[0]))
    drawing.remove(inline)
    drawing.append(anchor)

"""MKGDocument : moteur de generation Word a la charte MKG (couche Document V6).

Gere les sections (portrait/paysage), les pages << brand >> pleine page rendues
en Pillow (couverture, intercalaire, fin), l'en-tete / pied courant avec champ
de pagination Word, et l'ensemble des blocs de contenu natifs (titres, corps,
listes, tableaux denses, cartes KPI, encadres insight, sommaire, glossaire).
"""
from __future__ import annotations

import tempfile
import uuid
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Emu, Mm, Pt, RGBColor
from PIL import Image

from . import charte
from . import oxml_helpers as ox
from . import pillow_pages as pp
from . import sanitize

def _resolve_default_assets() -> Path:
    """Resout le dossier d'assets : d'abord embarque dans le package
    (mkg_docx/assets), puis repli sur l'arbo MKG/assets si presente."""
    bundled = Path(__file__).resolve().parent / "assets"
    if (bundled / "logo-mkg-bleu.png").exists() or (bundled / "logo-mkg-bleu.svg").exists():
        return bundled
    legacy = Path(__file__).resolve().parents[3] / "assets"
    if legacy.exists():
        return legacy
    return bundled


ASSETS_DEFAULT = _resolve_default_assets()


class MKGDocument:
    def __init__(self, *, orientation: str = "portrait", study_title: str = "",
                 date: str = "", confidential: str = charte.CONFIDENTIAL_LINE,
                 grain: bool = True, dpi: int = pp.DEFAULT_DPI,
                 assets_dir: str | Path | None = None):
        self.doc = Document()
        self.orientation = orientation.lower()
        self.study_title = study_title
        self.date = date
        self.confidential = confidential
        self.grain = grain
        self.dpi = dpi

        assets = Path(assets_dir) if assets_dir else ASSETS_DEFAULT
        # Logo bleu officiel (carre RVB, lettres blanches) : en-tete sur fond clair
        # et, place dans un cartouche blanc arrondi, sur les pages brand sombres.
        new_logo = assets / "logo-mkg-new-rvb.png"
        self.logo_blue = str(new_logo if new_logo.exists() else assets / "logo-mkg-bleu.png")
        self.logo_white = str(assets / "logo-mkg-blanc.png")

        self._mode: str | None = None
        self._tmp = Path(tempfile.mkdtemp(prefix="mkgdocx_"))
        self._grad_cache: dict[int, str] = {}

        self._configure_base_style()
        self._strip_initial_paragraph()

    # ------------------------------------------------------------------
    # Dimensions
    # ------------------------------------------------------------------
    @property
    def _page_size(self):
        return charte.A4_PORTRAIT if self.orientation == "portrait" else charte.A4_LANDSCAPE

    @property
    def _page_w(self):
        return self._page_size[0]

    @property
    def _page_h(self):
        return self._page_size[1]

    def _content_width_emu(self) -> int:
        return int(self._page_w) - 2 * int(charte.PAGE_MARGIN_X)

    def _content_width_mm(self) -> float:
        return (int(self._page_w) - 2 * int(charte.PAGE_MARGIN_X)) / 36000.0

    # ------------------------------------------------------------------
    # Base
    # ------------------------------------------------------------------
    def _configure_base_style(self) -> None:
        normal = self.doc.styles["Normal"]
        normal.font.name = charte.FONT
        normal.font.size = charte.SIZE_BODY
        normal.font.color.rgb = charte.rgb(charte.INK)
        pf = normal.paragraph_format
        pf.line_spacing = charte.LINE_BODY
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)

    def _strip_initial_paragraph(self) -> None:
        if self.doc.paragraphs:
            p0 = self.doc.paragraphs[0]
            p0._element.getparent().remove(p0._element)

    # ------------------------------------------------------------------
    # Runs / paragraphes stylises
    # ------------------------------------------------------------------
    def _run(self, para, text, *, weight="regular", size=charte.SIZE_BODY,
             color=charte.INK, italic=False, upper=False, tracking_em=0.0,
             raw=False):
        # Assainissement editorial (tirets / pictos) sauf glyphes de charte (raw).
        if not raw:
            text = sanitize.clean(text)
        if upper:
            text = text.upper()
        run = para.add_run(text)
        font = run.font
        name = {
            "regular": charte.FONT, "semibold": charte.FONT_SEMIBOLD,
            "bold": charte.FONT, "light": charte.FONT_LIGHT,
            "semilight": charte.FONT_SEMILIGHT,
        }.get(weight, charte.FONT)
        font.name = name
        ox.set_complex_font(run, name)
        font.size = size
        font.bold = weight == "bold"
        font.italic = italic
        font.color.rgb = charte.rgb(color)
        if tracking_em:
            ox.set_char_spacing(run, charte.tracking_twips(tracking_em, size.pt))
        return run

    def _p(self, *, space_before=0.0, space_after=0.0, line=None, align=None):
        self._ensure_content()
        para = self.doc.add_paragraph()
        pf = para.paragraph_format
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        if line is not None:
            pf.line_spacing = line
        if align is not None:
            para.alignment = align
        return para

    # ------------------------------------------------------------------
    # Sections
    # ------------------------------------------------------------------
    def _set_section_page(self, sec) -> None:
        if self.orientation == "portrait":
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = charte.A4_PORTRAIT
        else:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = charte.A4_LANDSCAPE

    def _next_section(self):
        if self._mode is None:
            return self.doc.sections[0]
        return self.doc.add_section(WD_SECTION.NEW_PAGE)

    def _new_brand_section(self):
        sec = self._next_section()
        self._set_section_page(sec)
        ox.make_section_fullbleed(sec)
        ox.unlink_headers_footers(sec)
        self._clear_hf(sec.header)
        self._clear_hf(sec.footer)
        self._mode = "brand"
        return sec

    def _ensure_content(self):
        if self._mode == "content":
            return self.doc.sections[-1]
        first = self._mode is None
        sec = self._next_section()
        self._set_section_page(sec)
        sec.top_margin = charte.PAGE_MARGIN_Y
        sec.bottom_margin = charte.PAGE_MARGIN_Y
        sec.left_margin = charte.PAGE_MARGIN_X
        sec.right_margin = charte.PAGE_MARGIN_X
        sec.gutter = Mm(0)
        sec.header_distance = charte.HEADER_DISTANCE
        sec.footer_distance = charte.FOOTER_DISTANCE
        ox.unlink_headers_footers(sec)
        if first:
            # document sans couverture : pas de pagination sur la 1re page
            sec.different_first_page_header_footer = True
            self._clear_hf(sec.first_page_header)
            self._clear_hf(sec.first_page_footer)
        self._build_runhead(sec)
        self._build_runfoot(sec)
        self._mode = "content"
        return sec

    @staticmethod
    def _clear_hf(part) -> None:
        for p in list(part.paragraphs):
            p._element.getparent().remove(p._element)
        part.add_paragraph()

    def _borderless_table(self, part_or_doc, rows, cols, width_emu):
        try:
            table = part_or_doc.add_table(rows=rows, cols=cols, width=Emu(width_emu))
        except TypeError:
            table = part_or_doc.add_table(rows=rows, cols=cols)
        table.autofit = False
        table.allow_autofit = False
        ox.set_table_borders(table, which=())
        # largeur explicite
        per = int(width_emu / cols)
        for row in table.rows:
            for cell in row.cells:
                cell.width = Emu(per)
        return table

    def _build_runhead(self, sec) -> None:
        hdr = sec.header
        self._clear_hf(hdr)
        # supprime le paragraphe vide initial puis pose une table 2 colonnes
        for p in list(hdr.paragraphs):
            p._element.getparent().remove(p._element)
        w = self._content_width_emu()
        table = self._borderless_table(hdr, 1, 2, w)
        left, right = table.rows[0].cells
        left.width = Emu(int(w * 0.8))
        right.width = Emu(int(w * 0.2))
        # filet bas
        ox.set_table_borders(table, color=charte.GRAY_100, size=4, which=("bottom",))
        lp = left.paragraphs[0]
        ox.set_no_spacing(lp)
        if self.study_title:
            self._run(lp, self.study_title, weight="semibold",
                      size=charte.SIZE_RUNHEAD, color=charte.INK_MUTED,
                      upper=True, tracking_em=charte.TRACK_RUNHEAD_EM)
        rp = right.paragraphs[0]
        ox.set_no_spacing(rp)
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.add_run().add_picture(self.logo_blue, height=Mm(4))

    def _build_runfoot(self, sec) -> None:
        ftr = sec.footer
        for p in list(ftr.paragraphs):
            p._element.getparent().remove(p._element)
        w = self._content_width_emu()
        table = self._borderless_table(ftr, 1, 3, w)
        third = int(w / 3)
        for c in table.rows[0].cells:
            c.width = Emu(third)
        ox.set_table_borders(table, color=charte.GRAY_100, size=4, which=("top",))
        c_left, c_mid, c_right = table.rows[0].cells
        lp = c_left.paragraphs[0]
        ox.set_no_spacing(lp)
        if self.date:
            self._run(lp, self.date, size=charte.SIZE_RUNFOOT, color=charte.INK_MUTED)
        mp = c_mid.paragraphs[0]
        ox.set_no_spacing(mp)
        mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if self.confidential:
            self._run(mp, self.confidential, weight="semibold",
                      size=charte.SIZE_RUNFOOT, color=charte.INK_MUTED,
                      tracking_em=0.08)
        rp = c_right.paragraphs[0]
        ox.set_no_spacing(rp)
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pg = ox.add_page_number_field(rp)
        pg.font.name = charte.FONT
        pg.font.bold = True
        pg.font.size = charte.SIZE_PAGENO
        pg.font.color.rgb = charte.rgb(charte.BLUE_DEEP)

    # ------------------------------------------------------------------
    # Pages BRAND (pleine page Pillow)
    # ------------------------------------------------------------------
    def _renderer(self):
        pw_mm = int(self._page_w) / 36000.0
        ph_mm = int(self._page_h) / 36000.0
        return pp.BrandPageRenderer(pw_mm, ph_mm, self.logo_white,
                                    logo_blue_path=self.logo_blue,
                                    dpi=self.dpi, grain=self.grain)

    def _place_brand_image(self, sec, img: Image.Image) -> None:
        path = self._tmp / f"brand_{uuid.uuid4().hex}.jpg"
        img.convert("RGB").save(path, "JPEG", quality=88, optimize=True)
        para = self.doc.add_paragraph()
        ox.set_no_spacing(para)
        para.add_run().font.size = Pt(1)
        ox.add_fullpage_background(para, str(path), int(self._page_w), int(self._page_h))

    def add_cover(self, *, title, eyebrow="", subtitle="", fields=None,
                  confidential="", address="", title_size_pt=32.0) -> None:
        sec = self._new_brand_section()
        spec = pp.CoverSpec(title=title, eyebrow=eyebrow, subtitle=subtitle,
                            fields=fields or [], confidential=confidential,
                            address=address, title_size_pt=title_size_pt)
        self._place_brand_image(sec, self._renderer().cover(spec))

    def add_divider(self, *, number, title, chapter_label="", items=None,
                    footer="") -> None:
        sec = self._new_brand_section()
        spec = pp.DividerSpec(number=str(number), title=title,
                              chapter_label=chapter_label, items=items or [],
                              footer=footer)
        self._place_brand_image(sec, self._renderer().divider(spec))

    def add_end(self, *, tagline=None, eyebrow=charte.SIGNATURE, fields=None,
                confidential="", date="") -> None:
        sec = self._new_brand_section()
        kwargs = dict(eyebrow=eyebrow, fields=fields or [],
                      confidential=confidential, date=date)
        if tagline:
            kwargs["tagline"] = tagline
        self._place_brand_image(sec, self._renderer().end(pp.EndSpec(**kwargs)))

    # ------------------------------------------------------------------
    # Blocs de contenu
    # ------------------------------------------------------------------
    def eyebrow(self, text: str) -> None:
        p = self._p(space_after=2)
        self._run(p, text, weight="semibold", size=charte.SIZE_EYEBROW,
                  color=charte.ELECTRIC, upper=True, tracking_em=charte.TRACK_EYEBROW_EM)

    def section_title(self, text: str, *, rule: bool = True,
                      space_before: float = 10.0) -> None:
        p = self._p(space_before=space_before, space_after=4, line=charte.LINE_H2)
        self._run(p, text, weight="bold", size=charte.SIZE_H2,
                  color=charte.BLUE_DEEP, upper=True, tracking_em=charte.TRACK_H2_EM)
        if rule:
            self.gradient_rule()

    def title_h1(self, text: str, *, space_before: float = 6.0) -> None:
        p = self._p(space_before=space_before, space_after=4)
        self._run(p, text, weight="bold", size=charte.SIZE_H1, color=charte.BLUE_DEEP)

    def subheading(self, text: str, *, space_before: float = 8.0) -> None:
        p = self._p(space_before=space_before, space_after=3, line=charte.LINE_H3)
        self._run(p, text, weight="semibold", size=charte.SIZE_H3, color=charte.INK)

    def paragraph(self, text: str, *, space_after: float = 6.0) -> None:
        p = self._p(space_after=space_after, line=charte.LINE_BODY)
        self._run(p, text, size=charte.SIZE_BODY, color=charte.INK)

    def small(self, text: str, *, space_after: float = 4.0) -> None:
        p = self._p(space_after=space_after)
        self._run(p, text, size=charte.SIZE_SMALL, color=charte.INK_SOFT)

    def bullets(self, items: list) -> None:
        """items : chaine simple, ou (amorce_gras, suite), ou {lead, text}."""
        for it in items:
            if isinstance(it, dict):
                lead, text = it.get("lead", ""), it.get("text", "")
            elif isinstance(it, (tuple, list)) and len(it) == 2:
                lead, text = it
            else:
                lead, text = "", str(it)
            p = self._p(space_after=3, line=charte.LINE_BODY)
            p.paragraph_format.left_indent = Mm(5)
            p.paragraph_format.first_line_indent = Mm(-3)
            self._run(p, "\u2014  ", weight="semibold", color=charte.BLUE_DEEP, raw=True)
            if lead:
                self._run(p, lead.rstrip() + " ", weight="semibold", color=charte.INK)
            self._run(p, text, size=charte.SIZE_BODY, color=charte.INK)

    def key_value(self, pairs: list, *, label_ratio: float = 0.30) -> None:
        """Tableau cle/valeur (fiche d'en-tete) : libelle a gauche sur tint,
        valeur a droite. Pas d'en-tete electric (ce n'est pas un tableau de donnees)."""
        self._ensure_content()
        pairs = [p for p in pairs if (p[0] or (len(p) > 1 and p[1]))]
        if not pairs:
            return
        w = self._content_width_emu()
        table = self.doc.add_table(rows=len(pairs), cols=2)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        ox.set_table_borders(table, color=charte.WHITE, size=8,
                             which=("top", "bottom", "left", "right", "insideH", "insideV"))
        lw = int(w * label_ratio)
        vw = w - lw
        for i, pair in enumerate(pairs):
            key = pair[0]
            val = pair[1] if len(pair) > 1 else ""
            kc, vc = table.rows[i].cells
            kc.width = Emu(lw)
            vc.width = Emu(vw)
            ox.set_cell_shading(kc, charte.TINT)
            ox.set_cell_shading(vc, charte.WHITE)
            ox.set_cell_margins(kc, top=90, bottom=90, left=140, right=140)
            ox.set_cell_margins(vc, top=90, bottom=90, left=140, right=140)
            ox.set_cell_vertical_alignment(kc, "center")
            kp = kc.paragraphs[0]
            ox.set_no_spacing(kp)
            self._run(kp, key, weight="semibold", size=charte.SIZE_BODY,
                      color=charte.BLUE_DEEP)
            vp = vc.paragraphs[0]
            ox.set_no_spacing(vp)
            self._run(vp, val, size=charte.SIZE_BODY, color=charte.INK)
        self.spacer(6)

    def gradient_rule(self) -> None:
        p = self._p(space_before=1, space_after=6)
        ox.set_no_spacing(p)
        w_emu = self._content_width_emu()
        path = self._gradient_path()
        p.add_run().add_picture(path, width=Emu(w_emu), height=Pt(1.4))

    def spacer(self, pt: float = 6.0) -> None:
        self._p(space_after=pt)

    def page_break(self) -> None:
        self._ensure_content()
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # -- sommaire --
    def summary(self, entries: list[dict]) -> None:
        """entries : [{number, title, desc}]"""
        for e in entries:
            p = self._p(space_after=7)
            num = e.get("number", "")
            if num:
                self._run(p, f"{num}  ", weight="bold", size=Pt(14),
                          color=charte.ELECTRIC)
            self._run(p, e.get("title", ""), weight="semibold", size=Pt(10),
                      color=charte.BLUE_DEEP)
            desc = e.get("desc")
            if desc:
                p.add_run().add_break()
                desc = sanitize.clean(desc)
                run = p.add_run("    " + desc) if num else p.add_run(desc)
                run.font.name = charte.FONT
                run.font.size = charte.SIZE_SMALL
                run.font.color.rgb = charte.rgb(charte.INK_SOFT)

    # -- glossaire --
    def glossary(self, entries: list[dict]) -> None:
        """entries : [{term, definition}]"""
        for e in entries:
            p = self._p(space_after=5, line=charte.LINE_BODY)
            self._run(p, e.get("term", ""), weight="bold", size=charte.SIZE_BODY,
                      color=charte.BLUE_DEEP)
            self._run(p, " \u2014 ", size=charte.SIZE_BODY, color=charte.INK, raw=True)
            self._run(p, e.get("definition", ""),
                      size=charte.SIZE_BODY, color=charte.INK)

    # -- encadre insight --
    def insight(self, text: str, *, eyebrow: str = "Insight cl\u00e9") -> None:
        self._ensure_content()
        w = self._content_width_emu()
        table = self.doc.add_table(rows=1, cols=1)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        cell.width = Emu(w)
        ox.set_table_borders(table, which=())
        ox.set_cell_shading(cell, charte.BLUE_DEEP)
        ox.set_cell_margins(cell, top=200, bottom=200, left=240, right=240)
        cp = cell.paragraphs[0]
        ox.set_no_spacing(cp)
        if eyebrow:
            self._run(cp, eyebrow, weight="semibold", size=charte.SIZE_EYEBROW,
                      color=charte.WHITE, upper=True, tracking_em=charte.TRACK_EYEBROW_EM)
            cp.add_run().add_break()
        r = cp.add_run(sanitize.clean(text))
        r.font.name = charte.FONT
        r.font.size = Pt(11)
        r.font.color.rgb = charte.rgb(charte.WHITE)
        self.spacer(6)

    # -- cartes KPI --
    def kpi_row(self, cards: list[dict]) -> None:
        """cards : [{label, value, trend?, trend_dir?(up/down/neutral), variant?(default/dark)}]"""
        self._ensure_content()
        n = len(cards)
        if n == 0:
            return
        w = self._content_width_emu()
        table = self.doc.add_table(rows=1, cols=n)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        ox.set_table_borders(table, color=charte.GRAY_100, size=4,
                             which=("top", "bottom", "left", "right", "insideV"))
        per = int(w / n)
        for i, card in enumerate(cards):
            cell = table.rows[0].cells[i]
            cell.width = Emu(per)
            ox.set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
            dark = card.get("variant") == "dark"
            if dark:
                ox.set_cell_shading(cell, charte.BLUE_DEEP)
            else:
                ox.set_cell_shading(cell, charte.WHITE)
            label_c = charte.WHITE if dark else charte.INK_SOFT
            value_c = charte.WHITE if dark else charte.BLUE_DEEP
            lp = cell.paragraphs[0]
            ox.set_no_spacing(lp)
            self._run(lp, card.get("label", ""), size=charte.SIZE_SMALL,
                      color=label_c, upper=False)
            vp = cell.add_paragraph()
            ox.set_no_spacing(vp)
            self._run(vp, card.get("value", ""), weight="bold", size=Pt(18),
                      color=value_c)
            trend = card.get("trend")
            if trend:
                tp = cell.add_paragraph()
                ox.set_no_spacing(tp)
                tdir = card.get("trend_dir", "neutral")
                if dark:
                    tcol = charte.DARK_POS if tdir == "up" else (
                        charte.DARK_NEG if tdir == "down" else charte.WHITE)
                else:
                    tcol = charte.POS if tdir == "up" else (
                        charte.NEG if tdir == "down" else charte.INK_SOFT)
                self._run(tp, trend, size=charte.SIZE_SMALL, color=tcol)
        self.spacer(6)

    # -- tableau dense --
    def table(self, headers: list[str], rows: list[list], *,
              footer: list | None = None, aligns: list[str] | None = None,
              col_ratios: list[float] | None = None, source: str = "") -> None:
        """aligns[i] dans {left,right,num,pos,neg}. col_ratios : largeurs relatives."""
        self._ensure_content()
        ncol = len(headers)
        aligns = aligns or ["left"] * ncol
        nrows = 1 + len(rows) + (1 if footer else 0)
        table = self.doc.add_table(rows=nrows, cols=ncol)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        ox.set_table_borders(table, color=charte.WHITE, size=4,
                             which=("top", "bottom", "insideH"))

        w = self._content_width_emu()
        if col_ratios and len(col_ratios) == ncol:
            total = sum(col_ratios)
            widths = [int(w * r / total) for r in col_ratios]
        else:
            widths = [int(w / ncol)] * ncol

        # en-tete
        head = table.rows[0]
        for j, htext in enumerate(headers):
            cell = head.cells[j]
            cell.width = Emu(widths[j])
            ox.set_cell_shading(cell, charte.ELECTRIC)
            ox.set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            ox.set_cell_vertical_alignment(cell, "center")
            p = cell.paragraphs[0]
            ox.set_no_spacing(p)
            p.alignment = self._align(aligns[j])
            self._run(p, htext, weight="semibold", size=charte.SIZE_TABLE_HEAD,
                      color=charte.WHITE, upper=True, tracking_em=0.06)
        ox.set_repeat_table_header(head)

        # corps
        for ri, rowdata in enumerate(rows):
            row = table.rows[1 + ri]
            shade = charte.TINT2 if ri % 2 == 0 else charte.TINT
            for j in range(ncol):
                cell = row.cells[j]
                cell.width = Emu(widths[j])
                ox.set_cell_shading(cell, shade)
                ox.set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
                ox.set_cell_vertical_alignment(cell, "center")
                p = cell.paragraphs[0]
                ox.set_no_spacing(p)
                p.alignment = self._align(aligns[j])
                val = "" if j >= len(rowdata) else str(rowdata[j])
                kind = aligns[j]
                if kind == "num":
                    self._run(p, val, weight="bold", size=charte.SIZE_TABLE_BODY,
                              color=charte.BLUE_DEEP)
                elif kind == "pos":
                    self._run(p, val, size=charte.SIZE_TABLE_BODY, color=charte.POS)
                elif kind == "neg":
                    self._run(p, val, size=charte.SIZE_TABLE_BODY, color=charte.NEG)
                else:
                    self._run(p, val, size=charte.SIZE_TABLE_BODY, color=charte.INK)

        # total
        if footer:
            frow = table.rows[-1]
            for j in range(ncol):
                cell = frow.cells[j]
                cell.width = Emu(widths[j])
                ox.set_cell_shading(cell, charte.BLUE_DEEP)
                ox.set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                p = cell.paragraphs[0]
                ox.set_no_spacing(p)
                p.alignment = self._align(aligns[j])
                val = "" if j >= len(footer) else str(footer[j])
                self._run(p, val, weight="bold", size=charte.SIZE_TABLE_FOOT,
                          color=charte.WHITE)

        if source:
            self.small(source, space_after=6)
        else:
            self.spacer(4)

    @staticmethod
    def _align(kind: str):
        if kind in ("right", "num", "pos", "neg"):
            return WD_ALIGN_PARAGRAPH.RIGHT
        if kind == "center":
            return WD_ALIGN_PARAGRAPH.CENTER
        return WD_ALIGN_PARAGRAPH.LEFT

    # ------------------------------------------------------------------
    # Gradient rule
    # ------------------------------------------------------------------
    def _gradient_path(self) -> str:
        w_px = max(200, int(self._content_width_mm() / 25.4 * 150))
        if w_px in self._grad_cache:
            return self._grad_cache[w_px]
        img = Image.new("RGBA", (w_px, 6), (0, 0, 0, 0))
        bd = charte.rgb_tuple(charte.BLUE_DEEP)
        el = charte.rgb_tuple(charte.ELECTRIC)
        px = img.load()
        for x in range(w_px):
            t = x / (w_px - 1)
            if t <= 0.5:
                k = t / 0.5
                col = tuple(round(bd[i] + (el[i] - bd[i]) * k) for i in range(3))
                a = 255
            else:
                k = (t - 0.5) / 0.5
                col = el
                a = round(255 * (1 - k))
            for y in range(6):
                px[x, y] = (col[0], col[1], col[2], a)
        path = self._tmp / "gradient_rule.png"
        img.save(path)
        self._grad_cache[w_px] = str(path)
        return str(path)

    # ------------------------------------------------------------------
    # Sauvegarde
    # ------------------------------------------------------------------
    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path

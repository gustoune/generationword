"""MKGDeck : moteur de generation PowerPoint a la charte MKG (deck 16:9).

Slides << brand >> (couverture, transition, stat, citation, merci, cloture)
rendues en image pleine page (Pillow), slides de contenu natives et editables
(bandeau bleu + corps + pied), construites a partir de blocs reutilisables
(paragraphes, KPI, tableau, liste numerotee, SWOT, recommandations, timeline,
takeaways).
"""
from __future__ import annotations

import tempfile
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

from . import charte, charts, icons, oxml, sanitize, units
from . import pillow_bg


def _resolve_assets() -> Path:
    bundled = Path(__file__).resolve().parent / "assets"
    if (bundled / "logo-mkg-blanc.png").exists():
        return bundled
    legacy = Path(__file__).resolve().parents[3] / "MKG" / "assets"
    return legacy if legacy.exists() else bundled


ASSETS_DEFAULT = _resolve_assets()


class MKGDeck:
    def __init__(self, *, doc_line: str = "", scale: int = 2, grain: bool = True,
                 assets_dir=None):
        self.prs = Presentation()
        self.prs.slide_width = units.emu(charte.W)
        self.prs.slide_height = units.emu(charte.H)
        self.doc_line = doc_line
        self.assets = Path(assets_dir) if assets_dir else ASSETS_DEFAULT
        new = self.assets / "logo-mkg-new-rvb.png"
        self.logo_footer = str(new if new.exists() else self.assets / "logo-mkg-bleu.png")
        self.logo_footer_white = str(self.assets / "logo-mkg-blanc.png")
        self._render = pillow_bg.BrandRenderer(self.assets, scale=scale, grain=grain)
        self._blank = self.prs.slide_layouts[6]
        self._tmp = Path(tempfile.mkdtemp(prefix="mkgppt_"))
        self._page = 0
        self._icon_cache = {}

    # ================================================================
    # Primitives
    # ================================================================
    def _slide(self):
        return self.prs.slides.add_slide(self._blank)

    def _rect(self, slide, x, y, w, h, *, fill=None, alpha=None, radius=None,
              line_hex=None, line_w=None, line_alpha=None):
        shape_t = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        sp = slide.shapes.add_shape(shape_t, units.emu(x), units.emu(y),
                                    units.emu(w), units.emu(h))
        sp.shadow.inherit = False
        if radius:
            oxml.corner_radius(sp, radius, min(w, h))
        if fill is None:
            sp.fill.background()
        else:
            oxml.shape_fill(sp, fill, alpha)
        if line_hex:
            oxml.line(sp, line_hex, units.pt(line_w or 1), line_alpha)
        else:
            oxml.no_line(sp)
        return sp

    def _oval(self, slide, x, y, d, *, fill=None, alpha=None, line_hex=None, line_w=None,
              line_alpha=None):
        sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, units.emu(x), units.emu(y),
                                    units.emu(d), units.emu(d))
        sp.shadow.inherit = False
        if fill is None:
            sp.fill.background()
        else:
            oxml.shape_fill(sp, fill, alpha)
        if line_hex:
            oxml.line(sp, line_hex, units.pt(line_w or 1), line_alpha)
        else:
            oxml.no_line(sp)
        return sp

    def _tf(self, slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
        tb = slide.shapes.add_textbox(units.emu(x), units.emu(y), units.emu(w), units.emu(h))
        oxml.tidy_textframe(tb.text_frame, anchor=anchor, wrap=wrap)
        return tb.text_frame

    def _para(self, tf, *, first=False, align=None, space_before=0, space_after=0,
              line=None):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        if space_before:
            p.space_before = units.pt(space_before)
        p.space_after = units.pt(space_after)
        if line is not None:
            p.line_spacing = line
        return p

    def _run(self, p, text, *, weight="regular", size=23, color=charte.INK, alpha=None,
             upper=False, tracking=None, italic=False):
        text = sanitize.clean(text)
        if upper and isinstance(text, str):
            text = text.upper()
        r = p.add_run()
        r.text = text or ""
        name, bold = charte.WEIGHTS.get(weight, (charte.FONT, False))
        r.font.name = name
        r.font.size = units.pt(size)
        r.font.bold = bold
        r.font.italic = italic
        oxml.run_color(r, color, alpha)
        if tracking:
            oxml.letter_spacing(r, tracking, size)
        return r

    def _simple(self, slide, x, y, w, h, text, **run_kw):
        anchor = run_kw.pop("anchor", MSO_ANCHOR.TOP)
        align = run_kw.pop("align", None)
        line = run_kw.pop("line", None)
        wrap = run_kw.pop("wrap", True)
        tf = self._tf(slide, x, y, w, h, anchor, wrap=wrap)
        p = self._para(tf, first=True, align=align, line=line)
        self._run(p, text, **run_kw)
        return tf

    def _picture(self, slide, path, x, y, *, w=None, h=None):
        if not Path(path).exists():
            return None
        kw = {}
        if w:
            kw["width"] = units.emu(w)
        if h:
            kw["height"] = units.emu(h)
        return slide.shapes.add_picture(str(path), units.emu(x), units.emu(y), **kw)

    # ---- icones ----
    def _icon_path(self, name, stroke_hex, width=2.0):
        key = (name, stroke_hex, round(width, 2))
        if key in self._icon_cache:
            return self._icon_cache[key]
        img = icons.render(name, stroke_hex, width=width)
        path = self._tmp / f"icon_{uuid.uuid4().hex}.png"
        img.save(path, "PNG")
        self._icon_cache[key] = str(path)
        return str(path)

    def _icon_badge(self, slide, x, y, size, *, name, bg, stroke,
                    shape="circle", radius=14, icon_frac=0.50, width=2.0):
        """Pastille (cercle ou carre arrondi) + pictogramme centre."""
        if shape == "circle":
            self._oval(slide, x, y, size, fill=bg)
        else:
            self._rect(slide, x, y, size, size, fill=bg, radius=radius)
        d = size * icon_frac
        off = (size - d) / 2
        if name:
            self._picture(slide, self._icon_path(name, stroke, width), x + off, y + off,
                          w=d, h=d)

    # ================================================================
    # Slides BRAND (image pleine page)
    # ================================================================
    def _full_image(self, img):
        path = self._tmp / f"brand_{uuid.uuid4().hex}.png"
        img.convert("RGB").save(path, "PNG")
        slide = self._slide()
        slide.shapes.add_picture(str(path), 0, 0, units.emu(charte.W), units.emu(charte.H))
        return slide

    def add_cover(self, *, title, eyebrow="", subtitle="", meta=None, kpis=None,
                  footer_note=charte.CONFIDENTIAL_LINE):
        self._full_image(self._render.cover(
            title=title, eyebrow=eyebrow, subtitle=subtitle, meta=meta or [],
            kpis=kpis or [], footer_note=footer_note))

    def add_section(self, *, number, title, eyebrow="", desc="", toc=None):
        """Intercalaire de chapitre : fond en image (charte fidele) mais TOUS les
        textes (numero geant, eyebrow, titre, description, sommaire) en zones
        natives editables et deplacables dans PowerPoint."""
        slide = self._slide()
        # 1. fond seul (sans texte) en image pleine page
        bg = self._render.section_bg()
        path = self._tmp / f"sectionbg_{uuid.uuid4().hex}.png"
        bg.convert("RGB").save(path, "PNG")
        slide.shapes.add_picture(str(path), 0, 0, units.emu(charte.W), units.emu(charte.H))

        px = 72
        split = 806
        # 2. numero geant decoratif (blanc tres faible), editable
        if str(number):
            self._simple(slide, px - 8, 262, 760, 340, str(number),
                         weight="black", size=200, color=charte.WHITE, alpha=5,
                         wrap=False)
        cy = int(charte.H * 0.40)
        # 3. eyebrow (libelle de chapitre)
        if eyebrow:
            self._simple(slide, px, cy, 700, 26, eyebrow,
                         weight="semibold", size=14, color=charte.WHITE, alpha=35,
                         upper=True, tracking=0.16, wrap=False)
        # 4. titre (panneau sombre gauche, wrap autorise)
        self._simple(slide, px, cy + 26, split - px - 24, 120, title,
                     weight="bold", size=72, color=charte.WHITE,
                     line=1.05, wrap=True)
        # 5. description
        if desc:
            self._simple(slide, px, cy + 26 + 104, 560, 220, desc,
                         weight="light", size=21, color=charte.WHITE, alpha=59,
                         line=1.5, wrap=True)
        # 6. sommaire a droite (pastilles + libelles natifs)
        if toc:
            ix = split + 70
            iy = int(charte.H * 0.40)
            dia = 38
            for i, item in enumerate(toc):
                active = isinstance(item, dict) and item.get("active")
                label = item.get("label") if isinstance(item, dict) else str(item)
                if active:
                    self._oval(slide, ix, iy, dia, fill=charte.WHITE)
                    self._simple(slide, ix, iy, dia, dia, str(i + 1),
                                 weight="semibold", size=15, color=charte.ELECTRIC,
                                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
                    self._simple(slide, ix + dia + 16, iy, 360, dia, label,
                                 weight="semibold", size=21, color=charte.WHITE,
                                 anchor=MSO_ANCHOR.MIDDLE, wrap=False)
                else:
                    self._oval(slide, ix, iy, dia, line_hex=charte.WHITE,
                               line_w=2, line_alpha=30)
                    self._simple(slide, ix, iy, dia, dia, str(i + 1),
                                 weight="regular", size=15, color=charte.WHITE, alpha=45,
                                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
                    self._simple(slide, ix + dia + 16, iy, 360, dia, label,
                                 weight="regular", size=21, color=charte.WHITE, alpha=45,
                                 anchor=MSO_ANCHOR.MIDDLE, wrap=False)
                iy += 60
        return slide

    def add_stat(self, *, number, unit="", eyebrow="", desc="", source="", dark=False):
        self._full_image(self._render.stat(
            number=number, unit=unit, eyebrow=eyebrow, desc=desc, source=source, dark=dark))

    def add_quote(self, *, text, author=""):
        self._full_image(self._render.quote(text=text, author=author))

    def add_merci(self, *, title="Merci", subtitle="", contact=""):
        self._full_image(self._render.merci(title=title, subtitle=subtitle, contact=contact))

    def add_closing(self, *, title, eyebrow="", subtitle="", contacts=None):
        self._full_image(self._render.closing(
            title=title, eyebrow=eyebrow, subtitle=subtitle, contacts=contacts or []))

    # ================================================================
    # Frame de contenu (bandeau + corps + pied)
    # ================================================================
    def _content_frame(self, *, eyebrow="", title="", badge="", paged=True):
        slide = self._slide()
        # bandeau bleu
        self._rect(slide, 0, 0, charte.W, charte.HEADER_H, fill=charte.BLUE)
        hx = charte.HEADER_PAD_X
        hw = charte.W - 2 * charte.HEADER_PAD_X - (320 if badge else 0)
        # eyebrow + titre en zones distinctes a position fixe (pas de rognage
        # vertical du a un ancrage centre multi-paragraphes ; pas de wrap : ce
        # sont des en-tetes mono-ligne comme dans la charte).
        if eyebrow:
            etf = self._tf(slide, hx, 26, hw, 26, anchor=MSO_ANCHOR.TOP, wrap=False)
            pe = self._para(etf, first=True)
            self._run(pe, eyebrow, weight="semibold", size=15, color=charte.WHITE,
                      alpha=48, upper=True, tracking=0.14)
            ttf = self._tf(slide, hx, 54, hw, 56, anchor=MSO_ANCHOR.TOP, wrap=False)
            pt = self._para(ttf, first=True)
        else:
            ttf = self._tf(slide, hx, 0, hw, charte.HEADER_H, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
            pt = self._para(ttf, first=True)
        self._run(pt, title, weight="bold", size=42, color=charte.WHITE)
        if badge:
            bw = 300
            bx = charte.W - charte.HEADER_PAD_X - bw
            by = (charte.HEADER_H - 46) / 2
            self._rect(slide, bx, by, bw, 46, fill=charte.WHITE, alpha=12, radius=6,
                       line_hex=charte.WHITE, line_w=1, line_alpha=20)
            self._simple(slide, bx, by, bw, 46, badge, weight="semibold", size=15,
                         color=charte.WHITE, upper=True, tracking=0.06,
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        # pied
        self._footer(slide, paged)
        body = (charte.BODY_PAD_X, charte.HEADER_H + charte.BODY_PAD_Y,
                charte.W - 2 * charte.BODY_PAD_X,
                charte.H - charte.HEADER_H - charte.FOOTER_H - 2 * charte.BODY_PAD_Y)
        return slide, body

    def _footer(self, slide, paged=True, *, dark=False, pad_x=None):
        pad_x = charte.HEADER_PAD_X if pad_x is None else pad_x
        fy = charte.H - charte.FOOTER_H
        if dark:
            self._rect(slide, 0, fy, charte.W, 1, fill=charte.WHITE, alpha=10)
            logo = self.logo_footer_white
            doc_col, doc_alpha = charte.WHITE, 40
            page_fill = charte.WHITE
            page_alpha = 15
        else:
            self._rect(slide, 0, fy, charte.W, charte.FOOTER_H, fill=charte.PANEL)
            self._rect(slide, 0, fy, charte.W, 1, fill=charte.LINE_LIGHT)
            logo = self.logo_footer
            doc_col, doc_alpha = charte.INK_FAINT, None
            page_fill, page_alpha = charte.BLUE, None
        self._picture(slide, logo, pad_x - 8, fy + (charte.FOOTER_H - 26) / 2, h=26)
        if self.doc_line:
            self._simple(slide, pad_x + 50, fy, 1200, charte.FOOTER_H,
                         self.doc_line, weight="regular", size=13, color=doc_col,
                         alpha=doc_alpha, upper=True, tracking=0.06, anchor=MSO_ANCHOR.MIDDLE)
        if paged:
            self._page += 1
            d = 36
            cx = charte.W - pad_x - d
            cy = fy + (charte.FOOTER_H - d) / 2
            self._oval(slide, cx, cy, d, fill=page_fill, alpha=page_alpha)
            self._simple(slide, cx, cy, d, d, f"{self._page:02d}", weight="bold", size=14,
                         color=charte.WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # ---- helpers de boite ----
    @staticmethod
    def _split2(box, gap=56):
        x, y, w, h = box
        cw = (w - gap) / 2
        return (x, y, cw, h), (x + cw + gap, y, cw, h)

    def _eyebrow(self, slide, box, text):
        """section-eyebrow ; renvoie le y suivant."""
        x, y, w, h = box
        self._simple(slide, x, y, w, 30, text, weight="semibold", size=17, color=charte.BLUE,
                     upper=True, tracking=0.12)
        return y + 42

    # ================================================================
    # BLOCS de corps
    # ================================================================
    def _block_paragraphs(self, slide, box, paras, *, eyebrow="", highlight=None):
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        tf = self._tf(slide, x, y, w, h - (y - box[1]))
        for i, para in enumerate(paras):
            p = self._para(tf, first=(i == 0), space_after=16, line=1.5)
            self._run(p, para, size=23, color=charte.INK)
        if highlight:
            hy = box[1] + box[3] - 150
            self._highlight(slide, (x, hy, w, 150), highlight)

    def _highlight(self, slide, box, hl):
        x, y, w, h = box
        self._rect(slide, x, y, w, h, fill=charte.BLUE, radius=10)
        tf = self._tf(slide, x + 28, y + 22, w - 56, h - 44, anchor=MSO_ANCHOR.MIDDLE)
        p = self._para(tf, first=True, space_after=8)
        self._run(p, hl.get("title", ""), weight="semibold", size=20, color=charte.WHITE)
        p2 = self._para(tf, line=1.5)
        self._run(p2, hl.get("text", ""), size=17, color=charte.WHITE, alpha=78)

    def _block_bullets(self, slide, box, items, *, eyebrow=""):
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        tf = self._tf(slide, x, y, w, h - (y - box[1]))
        for i, it in enumerate(items):
            lead, text = self._split_item(it)
            p = self._para(tf, first=(i == 0), space_after=12, line=1.45)
            self._run(p, "\u2014  ", weight="semibold", size=22, color=charte.BLUE)
            if lead:
                self._run(p, lead.rstrip() + " ", weight="semibold", size=22, color=charte.INK)
            self._run(p, text, size=22, color=charte.INK)

    @staticmethod
    def _split_item(it):
        if isinstance(it, dict):
            return it.get("lead", ""), it.get("text", "")
        if isinstance(it, (tuple, list)) and len(it) == 2:
            return it[0], it[1]
        return "", str(it)

    def _block_numbered(self, slide, box, items, *, eyebrow="", cols=2):
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        gap = 48
        cw = (w - gap * (cols - 1)) / cols
        rows = (len(items) + cols - 1) // cols
        rh = (box[1] + h - y) / rows
        for i, it in enumerate(items):
            lead, text = self._split_item(it)
            r, cdx = divmod(i, cols)
            ix = x + cdx * (cw + gap)
            iy = y + r * rh
            self._oval(slide, ix, iy, 44, fill=charte.BLUE_50)
            self._simple(slide, ix, iy, 44, 44, str(i + 1), weight="bold", size=18,
                         color=charte.BLUE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            tf = self._tf(slide, ix + 64, iy - 2, cw - 64, rh - 16, anchor=MSO_ANCHOR.TOP)
            p = self._para(tf, first=True, line=1.5)
            if lead:
                self._run(p, lead.rstrip() + " ", weight="semibold", size=20, color=charte.BLUE)
            self._run(p, text, size=20, color=charte.INK)

    def _block_kpis(self, slide, box, cards, *, eyebrow="", cols=None):
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        n = len(cards)
        cols = cols or (n if n <= 3 else (n + 1) // 2)
        rows = (n + cols - 1) // cols
        gap = 16
        cw = (w - gap * (cols - 1)) / cols
        ch = (box[1] + h - y - gap * (rows - 1)) / rows
        for i, c in enumerate(cards):
            r, cdx = divmod(i, cols)
            cx = x + cdx * (cw + gap)
            cy = y + r * (ch + gap)
            self._kpi_card(slide, (cx, cy, cw, ch), c)

    def _kpi_card(self, slide, box, c):
        x, y, w, h = box
        dark = c.get("variant") == "accent"
        self._rect(slide, x, y, w, h, fill=(charte.BLUE if dark else charte.PANEL), radius=10)
        pad = 28
        tf = self._tf(slide, x + pad, y + 24, w - 2 * pad, h - 40, anchor=MSO_ANCHOR.MIDDLE)
        p = self._para(tf, first=True, space_after=8)
        self._run(p, c.get("label", ""), size=15, color=(charte.WHITE if dark else charte.INK_FAINT),
                  alpha=(55 if dark else None), upper=True, tracking=0.05)
        pv = self._para(tf, space_after=6)
        self._run(pv, c.get("value", ""), weight="semibold", size=48,
                  color=(charte.WHITE if dark else charte.BLUE))
        trend = c.get("trend")
        if trend:
            tdir = c.get("trend_dir", "neutral")
            if dark:
                tcol, alpha = charte.WHITE, 60
            else:
                tcol = charte.SUCCESS if tdir == "up" else (
                    charte.DANGER if tdir == "down" else charte.INK_FAINT)
                alpha = None
            ptr = self._para(tf)
            self._run(ptr, trend, weight="regular", size=17, color=tcol, alpha=alpha)
        sub = c.get("sub")
        if sub:
            ps = self._para(tf)
            self._run(ps, sub, size=15, color=(charte.WHITE if dark else charte.INK_FAINT),
                      alpha=(50 if dark else None))

    def _block_table(self, slide, box, headers, rows, *, eyebrow="", footer=None,
                     aligns=None, widths=None):
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        ncol = len(headers)
        aligns = aligns or ["left"] * ncol
        nrows = 1 + len(rows) + (1 if footer else 0)
        gtable = slide.shapes.add_table(nrows, ncol, units.emu(x), units.emu(y),
                                        units.emu(w), units.emu(min(h, 60 * nrows))).table
        self._strip_table_style(gtable)
        gtable.first_row = False
        gtable.horz_banding = False
        # largeurs
        if widths and len(widths) == ncol:
            tot = sum(widths)
            for j in range(ncol):
                gtable.columns[j].width = units.emu(w * widths[j] / tot)
        else:
            for j in range(ncol):
                gtable.columns[j].width = units.emu(w / ncol)
        # en-tete
        for j, htext in enumerate(headers):
            self._cell(gtable.cell(0, j), htext, fill=charte.ELECTRIC, color=charte.WHITE,
                       weight="semibold", size=15, upper=True, tracking=0.06,
                       align=aligns[j])
        # corps
        for ri, row in enumerate(rows):
            shade = charte.ROW_ALT if ri % 2 == 0 else charte.WHITE
            for j in range(ncol):
                val = "" if j >= len(row) else str(row[j])
                kind = aligns[j]
                color, weight = charte.INK_SOFT, "regular"
                if kind == "key":
                    color, weight = charte.BLUE, "semibold"
                elif kind == "pos":
                    color, weight = charte.POS, "semibold"
                elif kind == "neg":
                    color, weight = charte.NEG, "semibold"
                self._cell(gtable.cell(1 + ri, j), val, fill=shade, color=color,
                           weight=weight, size=20, align=kind)
        # total
        if footer:
            for j in range(ncol):
                val = "" if j >= len(footer) else str(footer[j])
                self._cell(gtable.cell(nrows - 1, j), val, fill=charte.WHITE,
                           color=charte.BLUE, weight="semibold", size=20, align=aligns[j])

    def _cell(self, cell, text, *, fill, color, weight="regular", size=20, upper=False,
              tracking=None, align="left"):
        oxml.shape_fill(cell, fill)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = units.emu(20)
        cell.margin_right = units.emu(20)
        cell.margin_top = units.emu(8)
        cell.margin_bottom = units.emu(8)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT if align in ("right", "num", "pos", "neg") else (
            PP_ALIGN.CENTER if align == "center" else PP_ALIGN.LEFT)
        self._run(p, text, weight=weight, size=size, color=color, upper=upper,
                  tracking=tracking)

    @staticmethod
    def _strip_table_style(table):
        tbl = table._tbl
        tblPr = tbl.find(qn("a:tblPr"))
        if tblPr is not None:
            for sid in tblPr.findall(qn("a:tableStyleId")):
                tblPr.remove(sid)

    def _block_swot(self, slide, box, pros, cons, *, pro_title="Forces",
                    con_title="Points de vigilance"):
        left, right = self._split2(box, gap=24)
        self._swot_col(slide, left, pro_title, pros, charte.SWOT_PRO_BG,
                       charte.SWOT_PRO_BORDER, charte.SWOT_PRO_TITLE, charte.SUCCESS, "+")
        self._swot_col(slide, right, con_title, cons, charte.SWOT_CON_BG,
                       charte.SWOT_CON_BORDER, charte.SWOT_CON_TITLE, charte.DANGER, "-")

    def _swot_col(self, slide, box, title, items, bg, border, title_col, badge, mark):
        x, y, w, h = box
        self._rect(slide, x, y, w, h, fill=bg, radius=12, line_hex=border, line_w=1)
        pad = 44
        self._rect(slide, x + pad, y + 38, 50, 50, fill=badge, radius=11)
        self._simple(slide, x + pad, y + 38, 50, 50, mark, weight="bold", size=26,
                     color=charte.WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        self._simple(slide, x + pad + 66, y + 38, w - pad - 80, 50, title, weight="bold",
                     size=28, color=title_col, anchor=MSO_ANCHOR.MIDDLE)
        tf = self._tf(slide, x + pad, y + 120, w - 2 * pad, h - 150)
        for i, it in enumerate(items):
            p = self._para(tf, first=(i == 0), space_after=18, line=1.45)
            self._run(p, "-  ", weight="bold", size=21, color=title_col)
            self._run(p, str(it), size=21, color=charte.INK)

    def _block_reco(self, slide, box, items, *, eyebrow=""):
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        n = len(items)
        gap = 18
        ch = (box[1] + h - y - gap * (n - 1)) / n
        for i, it in enumerate(items):
            iy = y + i * (ch + gap)
            self._rect(slide, x, iy, w, ch, fill=charte.PANEL, radius=12)
            self._rect(slide, x, iy, 5, ch, fill=charte.BLUE)
            self._simple(slide, x + 32, iy, 70, ch, str(i + 1), weight="bold", size=40,
                         color=charte.BLUE, anchor=MSO_ANCHOR.MIDDLE)
            title = it.get("title", "") if isinstance(it, dict) else str(it)
            desc = it.get("desc", "") if isinstance(it, dict) else ""
            tf = self._tf(slide, x + 116, iy, w - 150, ch, anchor=MSO_ANCHOR.MIDDLE)
            p = self._para(tf, first=True, space_after=8)
            self._run(p, title, weight="semibold", size=24, color=charte.INK)
            if desc:
                pd = self._para(tf, line=1.5)
                self._run(pd, desc, size=19, color=charte.INK_MUTED)

    def _block_timeline(self, slide, box, phases, *, eyebrow=""):
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        n = len(phases)
        cw = w / n
        for i, ph in enumerate(phases):
            ix = x + i * cw
            d = 68
            self._oval(slide, ix, y, d, fill=(charte.BLUE if i % 2 else charte.BLUE_50))
            self._simple(slide, ix, y, d, d, str(i + 1), weight="bold", size=24,
                         color=(charte.WHITE if i % 2 else charte.BLUE),
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if i < n - 1:
                self._rect(slide, ix + d + 12, y + d / 2 - 1, cw - d - 24, 2, fill=charte.GRAY)
            tf = self._tf(slide, ix, y + d + 24, cw - 24, h - d - 40)
            if ph.get("year"):
                p = self._para(tf, first=True, space_after=8)
                self._run(p, ph["year"], weight="semibold", size=16, color=charte.ELECTRIC,
                          upper=True, tracking=0.08)
                pt = self._para(tf, space_after=10)
            else:
                pt = self._para(tf, first=True, space_after=10)
            self._run(pt, ph.get("title", ""), weight="semibold", size=23, color=charte.INK)
            if ph.get("desc"):
                pd = self._para(tf, line=1.5)
                self._run(pd, ph["desc"], size=17, color=charte.INK_MUTED)

    def _block_takeaways(self, slide, box, items, *, eyebrow=""):
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        n = len(items)
        gap = 24
        rh = (box[1] + h - y - gap * (n - 1)) / n
        for i, it in enumerate(items):
            iy = y + i * (rh + gap)
            m = min(64, rh)
            self._rect(slide, x, iy + (rh - m) / 2, m, m, fill=charte.BLUE_50, radius=14)
            self._simple(slide, x, iy + (rh - m) / 2, m, m, str(i + 1), weight="bold", size=26,
                         color=charte.BLUE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            lead, text = self._split_item(it)
            tf = self._tf(slide, x + m + 28, iy, w - m - 28, rh, anchor=MSO_ANCHOR.MIDDLE)
            p = self._para(tf, first=True, line=1.35)
            if lead:
                self._run(p, lead.rstrip() + " ", weight="bold", size=30, color=charte.BLUE)
            self._run(p, text, size=30, color=charte.INK)

    def _block_ranking(self, slide, box, items, *, eyebrow="", maxval=None):
        """Classement a barres : rang + nom + barre proportionnelle (top 3 bleu,
        4-6 electrique, reste gris) + valeur dans la barre."""
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        n = len(items)
        if n == 0:
            return

        def _num(it):
            v = it.get("value", 0) if isinstance(it, dict) else 0
            try:
                return float(v)
            except (TypeError, ValueError):
                return 0.0

        mx = maxval or max((_num(it) for it in items), default=1) or 1
        gap = 14
        slot = (box[1] + h - y - gap * (n - 1)) / n
        bar_h = min(28, slot)
        pos_w, name_w = 56, min(360, w * 0.30)
        track_x = x + pos_w + 24 + name_w + 20
        track_w = box[0] + w - track_x
        for i, it in enumerate(items):
            name = it.get("name", "") if isinstance(it, dict) else str(it)
            disp = it.get("display") if isinstance(it, dict) else None
            val = _num(it)
            iy = y + i * (slot + gap)
            cy = iy + (slot - bar_h) / 2
            pos_col = charte.BLUE if i < 3 else charte.GRAY
            fill = charte.BLUE if i < 3 else (charte.ELECTRIC if i < 6 else charte.GRAY)
            self._simple(slide, x, iy, pos_w, slot, f"{i + 1:02d}", weight="bold", size=22,
                         color=pos_col, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            self._simple(slide, x + pos_w + 24, iy, name_w, slot, name, weight="regular",
                         size=19, color=charte.INK, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
            self._rect(slide, track_x, cy, track_w, bar_h, fill=charte.PANEL, radius=6)
            fw = max(bar_h, track_w * (val / mx))
            self._rect(slide, track_x, cy, fw, bar_h, fill=fill, radius=6)
            if disp:
                self._simple(slide, track_x, cy, fw - 14, bar_h, disp, weight="bold",
                             size=14, color=charte.WHITE, align=PP_ALIGN.RIGHT,
                             anchor=MSO_ANCHOR.MIDDLE, wrap=False)

    def _block_linechart(self, slide, box, *, series, xlabels, eyebrow="", ymin=None,
                         ymax=None, yfmt="{:.0f}"):
        """Courbe temporelle : trace rasterise (Pillow) + legende native."""
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        avail_h = box[1] + h - y
        legend_h = 52 if any(s.get("label") for s in series) else 0
        chart_h = avail_h - legend_h
        img = charts.linechart(w, chart_h, series=series, xlabels=xlabels,
                               ymin=ymin, ymax=ymax, yfmt=yfmt)
        path = self._tmp / f"chart_{uuid.uuid4().hex}.png"
        img.save(path, "PNG")
        self._picture(slide, str(path), x, y, w=w, h=chart_h)
        if legend_h:
            ly = y + chart_h + 14
            cx = x
            palette = [charte.BLUE, charte.ELECTRIC, charte.GRAY]
            for si, sv in enumerate(series):
                lbl = sv.get("label")
                if not lbl:
                    continue
                col = sv.get("color") or palette[si % len(palette)]
                self._rect(slide, cx, ly + 14, 26, 5, fill=col, radius=2)
                tw = 16 + len(lbl) * 11
                self._simple(slide, cx + 36, ly, tw, 32, lbl, weight="regular", size=18,
                             color=charte.INK_SOFT, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
                cx += 36 + tw + 30

    def _block_photos(self, slide, box, items, *, eyebrow=""):
        """Trois (ou n) emplacements photo MKG : cadre sombre + icone + label,
        legende (titre + sous-titre) sous chaque emplacement."""
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        n = len(items)
        if n == 0:
            return
        gap = 26
        cw = (w - gap * (n - 1)) / n
        avail_h = box[1] + h - y
        cap_h = 96
        slot_h = avail_h - cap_h
        icon = self._icon_path("image", charte.WHITE, width=1.8)
        for i, it in enumerate(items):
            cx = x + i * (cw + gap)
            self._rect(slide, cx, y, cw, slot_h, fill=charte.DARK, radius=14)
            self._rect(slide, cx, y, cw, slot_h, fill=charte.ELECTRIC, alpha=14, radius=14)
            bd = 56
            bx = cx + (cw - bd) / 2
            by = y + slot_h / 2 - bd
            self._rect(slide, bx, by, bd, bd, fill=charte.WHITE, alpha=10, radius=12)
            self._picture(slide, icon, bx + bd * 0.27, by + bd * 0.27, w=bd * 0.46, h=bd * 0.46)
            label = it.get("label", "") if isinstance(it, dict) else str(it)
            self._simple(slide, cx + 16, by + bd + 14, cw - 32, 28, label, weight="semibold",
                         size=13, color=charte.WHITE, alpha=45, upper=True, tracking=0.10,
                         align=PP_ALIGN.CENTER)
            caption = it.get("caption", "") if isinstance(it, dict) else ""
            sub = it.get("sub", "") if isinstance(it, dict) else ""
            tf = self._tf(slide, cx, y + slot_h + 14, cw, cap_h - 14)
            if caption:
                pc = self._para(tf, first=True, space_after=6, line=1.2)
                self._run(pc, caption, weight="semibold", size=20, color=charte.DARK)
            if sub:
                psb = self._para(tf, first=(not caption), line=1.45)
                self._run(psb, sub, size=16, color=charte.INK_MUTED)

    def _block_table_grouped(self, slide, box, *, groups, subheaders, rows, eyebrow="",
                             label_header="", aligns=None):
        """Tableau a en-tetes groupes : bandeau de groupes (bleu) + sous-en-tetes
        (electrique) + corps zebre, classes de cellule key/pos/neg/lbl."""
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        ncol = 1 + sum(g.get("span", 1) for g in groups)
        aligns = aligns or ["center"] * (ncol - 1)
        nrows = 2 + len(rows)
        th = min(box[1] + h - y, 56 * nrows)
        gt = slide.shapes.add_table(nrows, ncol, units.emu(x), units.emu(y),
                                    units.emu(w), units.emu(th)).table
        self._strip_table_style(gt)
        gt.first_row = False
        gt.horz_banding = False
        lbl_w = max(220, w * 0.24)
        rest = (w - lbl_w) / (ncol - 1)
        gt.columns[0].width = units.emu(lbl_w)
        for j in range(1, ncol):
            gt.columns[j].width = units.emu(rest)
        # ligne 0 : groupes (la 1re colonne = libelle, fusionnee verticalement)
        gt.cell(0, 0).merge(gt.cell(1, 0))
        self._cell(gt.cell(0, 0), label_header, fill=charte.BLUE, color=charte.WHITE,
                   weight="bold", size=15, upper=True, tracking=0.04, align="left")
        col = 1
        for g in groups:
            span = g.get("span", 1)
            c0 = gt.cell(0, col)
            if span > 1:
                c0.merge(gt.cell(0, col + span - 1))
            self._cell(c0, g.get("label", ""), fill=charte.BLUE, color=charte.WHITE,
                       weight="bold", size=15, upper=True, tracking=0.04, align="center")
            col += span
        # ligne 1 : sous-en-tetes
        for j, sh in enumerate(subheaders):
            self._cell(gt.cell(1, 1 + j), sh, fill=charte.ELECTRIC, color=charte.WHITE,
                       weight="semibold", size=13, upper=True, tracking=0.03, align="center")
        # corps
        for ri, row in enumerate(rows):
            shade = charte.ROW_ALT if ri % 2 == 0 else charte.WHITE
            self._cell(gt.cell(2 + ri, 0), self._cell_text(row[0]), fill=shade,
                       color=charte.INK, weight="semibold", size=18, align="left")
            for j in range(1, ncol):
                raw = row[j] if j < len(row) else ""
                txt = self._cell_text(raw)
                cls = raw.get("cls") if isinstance(raw, dict) else aligns[j - 1]
                color, weight = charte.INK_SOFT, "regular"
                if cls == "key":
                    color, weight = charte.BLUE, "semibold"
                elif cls == "pos":
                    color, weight = charte.POS, "semibold"
                elif cls == "neg":
                    color, weight = charte.NEG, "semibold"
                self._cell(gt.cell(2 + ri, j), txt, fill=shade, color=color,
                           weight=weight, size=18, align="center")

    @staticmethod
    def _cell_text(v):
        if isinstance(v, dict):
            return str(v.get("v", ""))
        return str(v)

    # ================================================================
    # SLIDE avant / apres (comparaison plein ecran)
    # ================================================================
    def add_beforeafter(self, *, before, after, eyebrow="", title=""):
        """Slide comparaison : moitie gauche << avant >> (claire), moitie droite
        << apres >> (bleu + texture), pastille fleche centrale."""
        slide = self._slide()
        half = charte.W / 2
        self._rect(slide, 0, 0, half, charte.H, fill=charte.PANEL)
        panel = self._render.texture_panel(int(half), charte.H, charte.BLUE, opacity=0.12)
        ppath = self._tmp / f"bapanel_{uuid.uuid4().hex}.png"
        panel.save(ppath, "PNG")
        self._picture(slide, str(ppath), half, 0, w=half, h=charte.H)
        self._beforeafter_side(slide, 0, before, dark=False)
        self._beforeafter_side(slide, half, after, dark=True)
        # pastille fleche centrale
        d = 72
        self._oval(slide, half - d / 2, charte.H / 2 - d / 2, d, fill=charte.WHITE)
        ad = d * 0.42
        self._picture(slide, self._icon_path("arrow-right", charte.BLUE, width=2.4),
                      half - ad / 2, charte.H / 2 - ad / 2, w=ad, h=ad)
        return slide

    def _beforeafter_side(self, slide, ox, data, *, dark):
        pad = 84
        w = charte.W / 2 - 2 * pad
        tag = data.get("tag", "")
        headline = data.get("headline", "")
        points = data.get("points", [])
        # mesure verticale approximative pour centrer le bloc
        n_head = 2 if len(headline) > 26 else 1
        block_h = (60 if tag else 0) + n_head * 52 + 28 + len(points) * 58
        y = max(pad, (charte.H - block_h) / 2)
        if tag:
            tw = 60 + len(tag) * 11
            tag_bg = charte.WHITE if dark else charte.LINE_LIGHT
            tag_alpha = 18 if dark else None
            self._rect(slide, ox + pad, y, tw, 44, fill=tag_bg, alpha=tag_alpha, radius=22)
            self._simple(slide, ox + pad, y, tw, 44, tag, weight="semibold", size=16,
                         color=(charte.WHITE if dark else charte.INK_MUTED),
                         upper=True, tracking=0.12, align=PP_ALIGN.CENTER,
                         anchor=MSO_ANCHOR.MIDDLE, wrap=False)
            y += 60
        htf = self._tf(slide, ox + pad, y, w, n_head * 52 + 8)
        ph = self._para(htf, first=True, line=1.15)
        self._run(ph, headline, weight="bold", size=40,
                  color=(charte.WHITE if dark else charte.INK))
        y += n_head * 52 + 28
        mark = "check" if dark else "minus"
        mark_col = charte.WHITE if dark else charte.GRAY
        for pt in points:
            self._picture(slide, self._icon_path(mark, mark_col, width=2.2),
                          ox + pad, y + 4, w=26, h=26)
            tf = self._tf(slide, ox + pad + 42, y - 4, w - 42, 56)
            p = self._para(tf, first=True, line=1.4)
            self._run(p, str(pt), size=20,
                      color=(charte.WHITE if dark else charte.INK_SOFT),
                      alpha=(82 if dark else None))
            y += 58

    # ================================================================
    # API publique - slides GEN (fond clair/sombre, grand titre, grilles riches)
    # ================================================================
    def gen(self, *, eyebrow="", title="", subtitle="", body=None, columns=None,
            dark=False, title_size=46, intro=False):
        """Slide generique facon `slides.css .gen` : eyebrow + grand titre + corps.

        - `intro=True` : pas de corps, titre centre verticalement (slide d'accroche).
        - `body` : un bloc riche (features, cards, concept, steps, steps2, ...).
        """
        slide = self._slide()
        if dark:
            self._rect(slide, 0, 0, charte.W, charte.H, fill=charte.DARK)
        px = 100
        eb_col = charte.ELECTRIC if dark else charte.BLUE
        title_col = charte.WHITE if dark else charte.DARK
        sub_col = charte.WHITE if dark else charte.INK_MUTED
        sub_alpha = 62 if dark else None
        top = charte.H / 2 - 200 if intro else 64
        y = top
        if eyebrow:
            self._simple(slide, px, y, charte.W - 2 * px, 30, eyebrow, weight="semibold",
                         size=18, color=eb_col, upper=True, tracking=0.16, wrap=False)
            y += 44
        tsize = 76 if intro else title_size
        th = (tsize * 1.12) * (2 if len(title) > 46 else 1) + 10
        ttf = self._tf(slide, px, y, charte.W - 2 * px, th)
        pt = self._para(ttf, first=True, line=1.08)
        self._run(pt, title, weight="bold", size=tsize, color=title_col)
        y += th
        if subtitle:
            stf = self._tf(slide, px, y, min(charte.W - 2 * px, 1180), 90)
            ps = self._para(stf, first=True, line=1.5)
            self._run(ps, subtitle, size=24, color=sub_col, alpha=sub_alpha)
            y += 96
        self._footer(slide, paged=True, dark=dark, pad_x=px)
        if not intro and (body or columns):
            by = y + 24
            box = (px, by, charte.W - 2 * px, charte.H - charte.FOOTER_H - by - 28)
            if columns:
                boxes = [box] if len(columns) == 1 else self._split2(box)
                for blk, bx in zip(columns, boxes):
                    self._dispatch_block(slide, bx, blk, dark=dark)
            else:
                self._dispatch_block(slide, box, body, dark=dark)
        return slide

    # ================================================================
    # BLOCS RICHES (grilles de cartes facon .gen-grid*)
    # ================================================================
    def _block_features(self, slide, box, items, *, eyebrow=""):
        """Grille 4 box : carte bleu-clair, pastille ronde sombre + icone, h + d."""
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        cols = 2
        rows = (len(items) + cols - 1) // cols
        gap = 24
        cw = (w - gap) / cols
        ch = (h - gap * (rows - 1)) / rows
        for i, it in enumerate(items):
            r, cdx = divmod(i, cols)
            cx = x + cdx * (cw + gap)
            cy = y + r * (ch + gap)
            self._rect(slide, cx, cy, cw, ch, fill=charte.FEAT_BG, radius=14)
            d = 84
            iy = cy + (ch - d) / 2
            self._icon_badge(slide, cx + 40, iy, d, name=it.get("icon", "circle"),
                             bg=charte.DARK, stroke=charte.ELECTRIC, shape="circle",
                             icon_frac=0.46)
            tx = cx + 40 + d + 28
            tf = self._tf(slide, tx, cy, cw - (tx - cx) - 36, ch, anchor=MSO_ANCHOR.MIDDLE)
            p = self._para(tf, first=True, space_after=12, line=1.22)
            self._run(p, it.get("title", ""), weight="bold", size=27, color=charte.DARK)
            if it.get("text"):
                pd = self._para(tf, line=1.5)
                self._run(pd, it["text"], size=19, color=charte.INK_MUTED)

    def _block_cards(self, slide, box, items, *, eyebrow="", cols=None, variant="light"):
        """Grille 6 box. variant 'light' (bordure, icone au-dessus) ou 'tint'
        (fond bleu-clair, icone a gauche)."""
        x, y, w, h = box
        if eyebrow:
            y = self._eyebrow(slide, (x, y, w, h), eyebrow)
        n = len(items)
        cols = cols or (3 if n != 4 else 2)
        rows = (n + cols - 1) // cols
        gap = 22
        cw = (w - gap * (cols - 1)) / cols
        ch = (h - gap * (rows - 1)) / rows
        for i, it in enumerate(items):
            r, cdx = divmod(i, cols)
            cx = x + cdx * (cw + gap)
            cy = y + r * (ch + gap)
            if variant == "tint":
                self._rect(slide, cx, cy, cw, ch, fill=charte.BLUE_50, radius=14)
                d = 64
                self._icon_badge(slide, cx + 30, cy + 30, d, name=it.get("icon", "circle"),
                                 bg=charte.BLUE, stroke=charte.WHITE, shape="circle",
                                 icon_frac=0.48)
                tx = cx + 30 + d + 22
                tf = self._tf(slide, tx, cy + 24, cw - (tx - cx) - 26, ch - 40)
                p = self._para(tf, first=True, space_after=10, line=1.2)
                self._run(p, it.get("title", ""), weight="bold", size=23, color=charte.DARK)
                if it.get("text"):
                    pd = self._para(tf, line=1.5)
                    self._run(pd, it["text"], size=18, color=charte.INK_MUTED)
            else:
                self._rect(slide, cx, cy, cw, ch, fill=charte.WHITE, radius=14,
                           line_hex=charte.LINE_LIGHT, line_w=0.75)
                d = 66
                self._icon_badge(slide, cx + 32, cy + 28, d, name=it.get("icon", "circle"),
                                 bg=charte.BLUE_50, stroke=charte.BLUE, shape="circle",
                                 icon_frac=0.48)
                tf = self._tf(slide, cx + 32, cy + 28 + d + 18, cw - 64, ch - (28 + d + 18) - 20)
                p = self._para(tf, first=True, space_after=12, line=1.2)
                self._run(p, it.get("title", ""), weight="bold", size=24, color=charte.DARK)
                if it.get("text"):
                    pd = self._para(tf, line=1.5)
                    self._run(pd, it["text"], size=18, color=charte.INK_MUTED)

    def _block_concept(self, slide, box, items):
        """Cartes mecanisme : 2-3 cartes (grises / accent bleu) separees par des
        fleches rondes electriques. Chaque carte : icone + tag + titre + desc."""
        x, y, w, h = box
        n = len(items)
        if n == 0:
            return
        arrow = 72
        # cartes de largeur egale, fleches rondes entre elles
        unit = (w - (n - 1) * arrow) / n
        cx = x
        centers = []
        for i, it in enumerate(items):
            accent = bool(it.get("accent"))
            bg = charte.BLUE if accent else charte.PANEL
            self._rect(slide, cx, y, unit, h, fill=bg, radius=14)
            pad = 44
            d = 68
            ic = it.get("icon", "circle")
            if accent:
                self._rect(slide, cx + pad, y + pad, d, d, fill=charte.WHITE, alpha=15, radius=14)
                self._picture(slide, self._icon_path(ic, charte.WHITE),
                              cx + pad + d * 0.25, y + pad + d * 0.25, w=d * 0.5, h=d * 0.5)
            else:
                self._icon_badge(slide, cx + pad, y + pad, d, name=ic, bg=charte.WHITE,
                                 stroke=charte.BLUE, shape="round", radius=14, icon_frac=0.5)
            ty = y + pad + d + 26
            # tag tracke dans sa propre zone mono-ligne
            self._simple(slide, cx + pad, ty, unit - pad, 26, it.get("tag", ""),
                         weight="semibold", size=15,
                         color=(charte.WHITE if accent else charte.BLUE),
                         alpha=(65 if accent else None), upper=True, tracking=0.10)
            tf = self._tf(slide, cx + pad, ty + 40, unit - 2 * pad, h - (ty + 40 - y) - pad)
            ptitle = self._para(tf, first=True, space_after=18, line=1.15)
            self._run(ptitle, it.get("title", ""), weight="bold", size=30,
                      color=(charte.WHITE if accent else charte.DARK))
            if it.get("text"):
                pd = self._para(tf, line=1.55)
                self._run(pd, it["text"], size=18,
                          color=(charte.WHITE if accent else charte.INK_SOFT),
                          alpha=(80 if accent else None))
            centers.append(cx + unit)
            cx += unit + arrow
        # fleches rondes
        for i in range(n - 1):
            ax = centers[i] + (arrow - 56) / 2
            ay = y + h / 2 - 28
            self._oval(slide, ax, ay, 56, fill=charte.ELECTRIC)
            self._simple(slide, ax, ay, 56, 56, "\u203a", weight="bold", size=30,
                         color=charte.WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def _block_steps(self, slide, box, items, *, dark=False, flow=""):
        """Chaine d'etapes horizontale : cartes numerotees + chevrons + flux."""
        x, y, w, h = box
        if flow:
            h -= 56
        n = len(items)
        if n == 0:
            return
        arrow = 44
        cw = (w - (n - 1) * arrow) / n
        for i, it in enumerate(items):
            cx = x + i * (cw + arrow)
            if dark:
                self._rect(slide, cx, y, cw, h, fill=charte.WHITE, alpha=5, radius=14,
                           line_hex=charte.WHITE, line_w=1, line_alpha=10)
            else:
                self._rect(slide, cx, y, cw, h, fill=charte.PANEL, radius=14)
            d = 74
            self._oval(slide, cx + (cw - d) / 2, y + 34, d, fill=charte.ELECTRIC)
            self._simple(slide, cx + (cw - d) / 2, y + 34, d, d, str(i + 1), weight="bold",
                         size=30, color=charte.WHITE, align=PP_ALIGN.CENTER,
                         anchor=MSO_ANCHOR.MIDDLE)
            tf = self._tf(slide, cx + 20, y + 34 + d + 24, cw - 40, h - (34 + d + 24) - 26,
                          anchor=MSO_ANCHOR.TOP)
            p = self._para(tf, first=True, align=PP_ALIGN.CENTER, space_after=14, line=1.2)
            self._run(p, it.get("title", ""), weight="bold", size=23,
                      color=(charte.WHITE if dark else charte.DARK))
            if it.get("text"):
                pd = self._para(tf, align=PP_ALIGN.CENTER, line=1.45)
                self._run(pd, it["text"], size=16,
                          color=(charte.WHITE if dark else charte.INK_MUTED),
                          alpha=(60 if dark else None))
            if i < n - 1:
                achev = cx + cw + (arrow - 22) / 2
                self._simple(slide, achev, y + h / 2 - 18, 22, 36, "\u203a", weight="bold",
                             size=26, color=charte.ELECTRIC, align=PP_ALIGN.CENTER,
                             anchor=MSO_ANCHOR.MIDDLE)
        if flow:
            self._simple(slide, x, y + h + 18, w, 40, flow, weight="regular", size=19,
                         color=charte.ELECTRIC, italic=True)

    def _block_steps2(self, slide, box, items):
        """Etapes en cartes (claires) : pastille sombre + eyebrow + titre + desc."""
        x, y, w, h = box
        n = len(items)
        if n == 0:
            return
        gap = 20
        cw = (w - (n - 1) * gap) / n
        for i, it in enumerate(items):
            cx = x + i * (cw + gap)
            bg = charte.BLUE_50 if i % 2 == 0 else charte.PANEL
            self._rect(slide, cx, y, cw, h, fill=bg, radius=14)
            d = 66
            self._oval(slide, cx + (cw - d) / 2, y + 38, d, fill=charte.DARK)
            self._simple(slide, cx + (cw - d) / 2, y + 38, d, d, str(i + 1), weight="bold",
                         size=26, color=charte.WHITE, align=PP_ALIGN.CENTER,
                         anchor=MSO_ANCHOR.MIDDLE)
            tf = self._tf(slide, cx + 24, y + 38 + d + 24, cw - 48, h - (38 + d + 24) - 28)
            pe = self._para(tf, first=True, align=PP_ALIGN.CENTER, space_after=12)
            self._run(pe, it.get("eyebrow", ""), weight="semibold", size=13, color=charte.BLUE,
                      upper=True, tracking=0.10)
            ph = self._para(tf, align=PP_ALIGN.CENTER, space_after=14, line=1.2)
            self._run(ph, it.get("title", ""), weight="bold", size=25, color=charte.DARK)
            if it.get("text"):
                pd = self._para(tf, align=PP_ALIGN.CENTER, line=1.5)
                self._run(pd, it["text"], size=17, color=charte.INK_MUTED)

    # ================================================================
    # API publique - slides de contenu
    # ================================================================
    def content(self, *, eyebrow="", title="", badge="", body=None, columns=None):
        """Slide de contenu generique.

        - `body` : un bloc unique (dict {type, ...}).
        - `columns` : liste de 1 a 2 blocs disposes en colonnes.
        """
        slide, box = self._content_frame(eyebrow=eyebrow, title=title, badge=badge)
        if columns:
            boxes = [box] if len(columns) == 1 else self._split2(box)
            for blk, bx in zip(columns, boxes):
                self._dispatch_block(slide, bx, blk)
        elif body:
            self._dispatch_block(slide, box, body)
        return slide

    def _dispatch_block(self, slide, box, blk, *, dark=False):
        t = blk.get("type")
        eb = blk.get("eyebrow", "")
        if t == "features":
            self._block_features(slide, box, blk.get("items", []), eyebrow=eb)
            return
        if t == "cards":
            self._block_cards(slide, box, blk.get("items", []), eyebrow=eb,
                              cols=blk.get("cols"), variant=blk.get("variant", "light"))
            return
        if t == "concept":
            self._block_concept(slide, box, blk.get("items", []))
            return
        if t == "steps":
            self._block_steps(slide, box, blk.get("items", []), dark=dark,
                              flow=blk.get("flow", ""))
            return
        if t == "steps2":
            self._block_steps2(slide, box, blk.get("items", []))
            return
        if t == "ranking":
            self._block_ranking(slide, box, blk.get("items", []), eyebrow=eb,
                                maxval=blk.get("maxval"))
            return
        if t == "linechart":
            self._block_linechart(slide, box, series=blk.get("series", []),
                                  xlabels=blk.get("xlabels", []), eyebrow=eb,
                                  ymin=blk.get("ymin"), ymax=blk.get("ymax"),
                                  yfmt=blk.get("yfmt", "{:.0f}"))
            return
        if t == "photos":
            self._block_photos(slide, box, blk.get("items", []), eyebrow=eb)
            return
        if t == "table_grouped":
            self._block_table_grouped(slide, box, groups=blk.get("groups", []),
                                      subheaders=blk.get("subheaders", []),
                                      rows=blk.get("rows", []), eyebrow=eb,
                                      label_header=blk.get("label_header", ""),
                                      aligns=blk.get("aligns"))
            return
        if t == "paragraphs":
            self._block_paragraphs(slide, box, blk.get("paras", []), eyebrow=eb,
                                   highlight=blk.get("highlight"))
        elif t == "bullets":
            self._block_bullets(slide, box, blk.get("items", []), eyebrow=eb)
        elif t == "numbered":
            self._block_numbered(slide, box, blk.get("items", []), eyebrow=eb,
                                 cols=blk.get("cols", 2))
        elif t == "kpis":
            self._block_kpis(slide, box, blk.get("cards", []), eyebrow=eb, cols=blk.get("cols"))
        elif t == "table":
            self._block_table(slide, box, blk.get("headers", []), blk.get("rows", []),
                              eyebrow=eb, footer=blk.get("footer"), aligns=blk.get("aligns"),
                              widths=blk.get("widths"))
        elif t == "swot":
            self._block_swot(slide, box, blk.get("pros", []), blk.get("cons", []),
                            pro_title=blk.get("pro_title", "Forces"),
                            con_title=blk.get("con_title", "Points de vigilance"))
        elif t == "reco":
            self._block_reco(slide, box, blk.get("items", []), eyebrow=eb)
        elif t == "timeline":
            self._block_timeline(slide, box, blk.get("phases", []), eyebrow=eb)
        elif t == "takeaways":
            self._block_takeaways(slide, box, blk.get("items", []), eyebrow=eb)

    def add_agenda(self, *, title="Sommaire", eyebrow="", subtitle="", items=None):
        slide = self._slide()
        split = int(charte.W * 0.40)
        self._rect(slide, 0, 0, split, charte.H, fill=charte.BLUE)
        tf = self._tf(slide, 64, 0, split - 128, charte.H, anchor=MSO_ANCHOR.MIDDLE)
        if eyebrow:
            p = self._para(tf, first=True, space_after=16)
            self._run(p, eyebrow, weight="semibold", size=14, color=charte.WHITE, alpha=40,
                      upper=True, tracking=0.16)
            pt = self._para(tf, space_after=16)
        else:
            pt = self._para(tf, first=True, space_after=16)
        self._run(pt, title, weight="bold", size=64, color=charte.WHITE)
        if subtitle:
            ps = self._para(tf, line=1.4)
            self._run(ps, subtitle, weight="light", size=18, color=charte.WHITE, alpha=55)
        # items a droite
        items = items or []
        rx = split + 72
        rw = charte.W - rx - 72
        ry = 120
        rh = (charte.H - 240) / max(1, len(items))
        for i, it in enumerate(items):
            iy = ry + i * rh
            self._oval(slide, rx, iy, 46, fill=charte.BLUE)
            self._simple(slide, rx, iy, 46, 46, f"{i + 1:02d}", weight="bold", size=18,
                         color=charte.WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            title_i = it.get("title", "") if isinstance(it, dict) else str(it)
            desc_i = it.get("desc", "") if isinstance(it, dict) else ""
            tfb = self._tf(slide, rx + 70, iy, rw - 70, rh - 10, anchor=MSO_ANCHOR.MIDDLE)
            p = self._para(tfb, first=True, space_after=5)
            self._run(p, title_i, weight="semibold", size=26, color=charte.INK)
            if desc_i:
                pd = self._para(tfb, line=1.5)
                self._run(pd, desc_i, size=18, color=charte.INK_MUTED)
            if i < len(items) - 1:
                self._rect(slide, rx, iy + rh - 6, rw, 1, fill=charte.LINE_LIGHT)
        return slide

    # ================================================================
    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path
